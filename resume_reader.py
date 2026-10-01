import re
import io
import pdfplumber
import docx
def extract_text_from_pdf(file_bytes: bytes) -> str:
    text_chunks = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_chunks.append(page_text)
    return "\n".join(text_chunks)
def extract_text_from_docx(file_bytes: bytes) -> str:
    document = docx.Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)
    return "\n".join(paragraphs)
def clean_text(text: str) -> str:
    if not text:
        return ""
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n", text)
    text = re.sub(r"[•●▪◦‣]", "-", text)
    return text.strip()
def read_resume(filename: str, file_bytes: bytes) -> str:
    filename_lower = filename.lower()

    if filename_lower.endswith(".pdf"):
        raw_text = extract_text_from_pdf(file_bytes)
    elif filename_lower.endswith(".docx"):
        raw_text = extract_text_from_docx(file_bytes)
    else:
        raise ValueError(f"Unsupported file type: {filename}")

    return clean_text(raw_text)


def truncate_for_display(text: str, max_chars: int = 500) -> str:
    if len(text) <= max_chars:
        return text
    return text[:max_chars].rstrip() + "..."
if __name__ == "__main__":
    import sys
    if len(sys.argv) > 1:
        path = sys.argv[1]
        with open(path, "rb") as f:
            content = f.read()
        print(read_resume(path, content)[:1000])
    else:
        print("Usage: python resume_reader.py <path-to-resume-file>")
